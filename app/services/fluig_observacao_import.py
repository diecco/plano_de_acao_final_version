import hashlib
import json
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import load_workbook


ITEM_CODES = tuple(
    [f"A{i:02d}" for i in range(1, 7)]
    + [f"B{i:02d}" for i in range(1, 10)]
    + [f"C{i:02d}" for i in range(1, 12)]
    + ["D01", "D02"]
    + [f"E{i:02d}" for i in range(1, 6)]
    + [f"F{i:02d}" for i in range(1, 4)]
)


class FluigImportError(ValueError):
    pass


def _text(value):
    if value is None:
        return ""
    return str(value).strip()


def _identifier(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = _text(value)
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def _find(headers, *candidates):
    normalized = {str(value).strip().casefold(): value for value in headers}
    for candidate in candidates:
        found = normalized.get(candidate.casefold())
        if found is not None:
            return found
    return None


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise FluigImportError(f"Data de observação inválida: {text or 'não informada'}")


def _parse_time(value):
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    text = _text(value)
    for pattern in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, pattern).time()
        except ValueError:
            pass
    return time(0, 0)


def _count(value):
    text = _text(value)
    if not text:
        return 0
    try:
        number = int(float(text.replace(",", ".")))
    except ValueError as exc:
        raise FluigImportError(f"Quantidade inválida: {text}") from exc
    if number < 0 or number > 999:
        raise FluigImportError(f"Quantidade fora do intervalo permitido: {text}")
    return number


def parse_fluig_observacoes(file_bytes):
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise FluigImportError("O arquivo não é uma planilha Excel válida.") from exc
    if not workbook.worksheets:
        raise FluigImportError("A planilha não possui abas.")

    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration as exc:
        raise FluigImportError("A planilha está vazia.") from exc

    solicitacao_col = _find(headers, "Solicitação")
    calendario_col = _find(headers, "Calendario - 001")
    if not solicitacao_col or not calendario_col or not _find(headers, "A1 - 001"):
        raise FluigImportError(
            "O arquivo não corresponde à exportação de Observação Comportamental do Fluig."
        )

    index = {header: position for position, header in enumerate(headers)}

    def value(row, *names):
        column = _find(headers, *names)
        return row[index[column]] if column else None

    registros = []
    erros = []
    solicitacoes = set()
    for excel_row, row in enumerate(rows, start=2):
        if not any(cell is not None and _text(cell) for cell in row):
            continue
        try:
            solicitacao = _identifier(value(row, "Solicitação"))
            if not solicitacao:
                raise FluigImportError("Solicitação não informada")
            if solicitacao in solicitacoes:
                raise FluigImportError("Solicitação duplicada no arquivo")
            solicitacoes.add(solicitacao)

            inicio = value(row, "Início")
            data_observacao = _parse_date(value(row, "Calendario - 001"))
            respostas = {}
            for code in ITEM_CODES:
                fluig_code = f"{code[0]}{int(code[1:])} - 001"
                quantidade = _count(value(row, fluig_code))
                if quantidade:
                    respostas[code] = quantidade

            situacao = _text(value(row, "Situação"))
            pessoas = _count(value(row, "Num_Obs - 001")) or 1
            registro = {
                "solicitacao": solicitacao,
                "data_observacao": data_observacao,
                "hora_observacao": _parse_time(inicio),
                "observador_nome": _text(value(row, "Nome_Lancamento - 001", "Requisitante")),
                "observador_matricula": _text(value(row, "Matricula  Requisitante")),
                "centro_custos_codigo": _text(value(row, "Cod_CC - 001")),
                "centro_custos_descricao": _text(value(row, "Nome_CC - 001")),
                "local_observado": _text(value(row, "Contrato_Obs - 001", "Nome_CC - 001")),
                "setor_observado": _text(value(row, "Setor_OBS - 001")) or "Não informado no Fluig",
                "area": _text(value(row, "radio_types - 001")) or "Não informada",
                "atividade": _text(value(row, "Contrato_Func - 001")),
                "pessoas_observadas": pessoas,
                "status": "concluida" if situacao.casefold() == "finalizada" else "rascunho",
                "situacao_fluig": situacao,
                "respostas": respostas,
            }
            payload = json.dumps(registro, ensure_ascii=False, default=str, sort_keys=True)
            registro["conteudo_hash"] = hashlib.sha256(payload.encode("utf-8")).hexdigest()
            registros.append(registro)
        except FluigImportError as exc:
            erros.append({"linha": excel_row, "erro": str(exc)})

    if not registros:
        raise FluigImportError("Nenhum registro válido foi encontrado na planilha.")
    return registros, erros
